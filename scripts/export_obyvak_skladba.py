#!/usr/bin/env python3
"""Export Czech material-schedule workbook for Obývák šikminy / kastlík / štíty.

Quantities from the current `models/obyvak` geometry. StoSilent masses use the
published application-guideline rates (Top Basic 2.5 kg/m², Top Finish 3.0 kg/m²).
Unit prices are retail snapshots (Kč s DPH) with source notes — Sto is quote-only.

    source .venv/bin/activate
    PYTHONPATH=models:src python scripts/export_obyvak_skladba.py
"""

from __future__ import annotations

import csv
import math
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "models"))
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from obyvak import build
from obyvak_geom import (
    LABEL_BASS_CD,
    LABEL_BASS_GKB,
    LABEL_BASS_HANGER,
    LABEL_BASS_WOOL,
    LABEL_PLENUM_WOOL,
    LABEL_RACKING_STRAP,
    LABEL_RAFTERS,
    LABEL_SLOPE_CD,
    LABEL_SLOPE_DIRECT,
    LABEL_SLOPE_NONIUS,
    LABEL_SOFFIT_CD,
    LABEL_SOFFIT_UD,
    LABEL_SOFFIT_DUCT,
    LABEL_SOFFIT_NONIUS,
    LABEL_SOFFIT_WOOL,
    ObyvakLayout,
    ObyvakParams,
)

OUT_XLSX = ROOT / "docs" / "obyvak-skladba-materialu.xlsx"
OUT_CSV = ROOT / "docs" / "obyvak-skladba-souhrn.csv"
OUT_HTML = ROOT / "docs" / "skladba.html"

# TL / application guideline (StoSilent Distance checklist + TDS Finish):
# 1st coat Top Basic ≈ 2.5 kg/m²; 2nd coat Top Finish ≈ 3.0 kg/m².
STOSILENT_BASIC_KG_M2 = 2.5
STOSILENT_FINISH_KG_M2 = 3.0
STOSILENT_PAIL_KG = 18.0

URL_STOSILENT_FINISH = "https://www.sto.cz/s/p/a1F2p00000Piv9kEAB/stosilent-top-finish"
URL_STOSILENT_BASIC = "https://www.sto.cz/s/p/a1F2p00000Piv9jEAB/stosilent-top-basic"
URL_STOSILENT_GUIDE = (
    "https://stoprod.e-spirit.cloud/cepcom/no/documents/Systemdokumentasjon/"
    "Akustikk/StoSilentDistance/Arbeidsanvisning/StoSilent-Distance-Application-guideline.pdf"
)
URL_STOSILENT_FINISH_TDS = (
    "https://www.sto.co.nz/00_common/tds/StoSilent%20TDS/StoSilent%20Top%20Finish%20TDS.PDF"
)
URL_NATURHELD_140 = "https://naturheld-izolace.cz/produkt/naturheld-140/"
URL_NATURHELD_80 = (
    "https://www.prirodnistavba.cz/drevovlaknita-deska-natur-held-140-kg-m3-"
    "1880-615-mm-tl-80-mm-pero-drazka-15725.html"
)
URL_NATURHELD_40 = (
    "https://www.prirodnistavba.cz/drevovlaknita-deska-natur-held-140-kg-m3-"
    "1250-600-mm-tl-40-mm-15722.html"
)
URL_RF = "https://www.rigips.cz/produkty/protipozarni-deska-rf-df/"
URL_RF_DEK = (
    "https://www.dek.cz/produkty/detail/3630042500-sadrokarton-po-deska-rf-12-5mm-1250-2000mm"
)
URL_JUTAFOL = "https://www.juta.cz/cs/produkty/parozabrany"
URL_JUTAFOL_SHOP = (
    "https://www.izomat.cz/parozabrana-juta-jutafol-n-al-170-special.html"
)
URL_GKB = "https://www.rigips.cz/produkty/stavebni-deska-rb-a/"
URL_CD = "https://www.rigips.cz/produkty/r-cd-profil/"
URL_CD_DEK = "https://www.dek.cz/produkty/vypis/4167-cd-profily-na-strop/28565-rigips"
URL_DIRECT = "https://www.rigips.cz/produkty/primy-zaves/"
URL_DIRECT_DEK = (
    "https://www.dek.cz/produkty/detail/3630551425-primy-zaves-60-125-mm-1-00mm-100-ks-bal"
)
URL_NONIUS_TOP = "https://www.rigips.cz/produkty/nonius-horni-dil/"
URL_NONIUS_TOP_SHOP = "https://www.izomat.cz/zaves-rigips-nonius-horni-340-mm-100-ks-baleni.html"
URL_NONIUS_BOT = "https://www.rigips.cz/produkty/nonius-spodni-dil-pro-cd/"
URL_NONIUS_BOT_DEK = "https://www.dek.cz/produkty/detail/3631100745-nonius-spodni-cd-dkm"
URL_NONIUS_PIN = "https://www.rigips.cz/produkty/nonius-pojistna-zavlacka/"
URL_NONIUS_PIN_SHOP = "https://www.izomat.cz/pojistna-zavlacka-nonius-dk-mont-100-ks-baleni.html"
URL_TRMEN = "https://www.rigips.cz/produkty/akusticky-trmen-se-sylomerem/"

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
SECTION_FILL = PatternFill("solid", fgColor="D6E3F0")
SECTION_FONT = Font(bold=True, name="Calibri", size=11)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
OPEN_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
MONEY = Alignment(wrap_text=True, vertical="top", horizontal="right")

COLS = (
    "Pořadí",
    "Vrstva / prvek",
    "Typ výrobku",
    "Zvolený výrobek (odkaz)",
    "Tloušťka / rozměr",
    "Množství netto",
    "Jednotka",
    "Doporučená rezerva",
    "Množství s rezervou",
    "Jedn. cena (Kč s DPH)",
    "Cena s rezervou (Kč)",
    "Zdroj ceny / TL",
    "Poznámka",
)


@dataclass
class Line:
    order: str | int
    layer: str
    product_type: str
    product: str  # "Name\nURL" or plain
    dimension: str
    qty_net: float
    unit: str
    reserve_pct: float | None  # None = n/a
    unit_price: float | None  # Kč s DPH per unit of qty
    price_source: str
    note: str
    open_item: bool = False

    @property
    def qty_with_reserve(self) -> float:
        if self.reserve_pct is None:
            return self.qty_net
        return self.qty_net * (1.0 + self.reserve_pct / 100.0)

    @property
    def line_total(self) -> float | None:
        if self.unit_price is None:
            return None
        return self.qty_with_reserve * self.unit_price


def _fmt(value: float, digits: int = 1) -> float | int:
    if isinstance(value, (int, float)) and abs(value - round(value)) < 1e-9:
        return int(round(value))
    return round(float(value), digits)


def _ceil_pack(qty: float, pack: float) -> int:
    return int(math.ceil(qty / pack - 1e-12))


def _longest(part) -> float:
    bb = part.bounding_box()
    return max(bb.size.X, bb.size.Y, bb.size.Z)


def _volume_m3(parts) -> float:
    return sum(float(p.volume) for p in parts) / 1e9


def _length_m(parts) -> float:
    return sum(_longest(p) for p in parts) / 1000.0


def _product(name: str, url: str = "") -> str:
    return f"{name}\n{url}" if url else name


def collect_quantities():
    p = ObyvakParams()
    g = ObyvakLayout(p)
    shape, _meta = build(p)
    by = defaultdict(list)
    for child in shape.children:
        by[child.label].append(child)

    slope_face = (g.x_ridge / g.cos + (g.x_furn - g.x_ridge) / g.cos) * p.room_length / 1e6
    trap_zs = [g.z_slope_offset(x, 0.0) for x in range(0, int(p.room_width) + 1, 50)]
    trap_h = sum(z - p.predstena_bottom_z for z in trap_zs) / len(trap_zs)
    trap_face = trap_h * p.room_width / 1e6
    soffit_y = g.y_furn1 - g.y_furn0
    nh_pts = g.soffit_nh_pts()
    bottom = abs(nh_pts[1][0] - nh_pts[0][0])
    vert = abs(nh_pts[4][1] - nh_pts[3][1])
    soffit_face = (bottom + vert) * soffit_y / 1e6
    lid_area = max(g.poz_r0 - g.x_gkf_kink, 0.0) * soffit_y / 1e6

    nonius_hang = []
    drops = []
    wall_brackets = []
    joint_angles = []
    cleats = []
    for part in by[LABEL_SOFFIT_NONIUS]:
        bb = part.bounding_box()
        sx, sy, sz = bb.size.X, bb.size.Y, bb.size.Z
        long = max(sx, sy, sz)
        if abs(sx - 20) < 1 and abs(sy - 20) < 1 and sz > 50:
            nonius_hang.append(part)
        elif abs(sx - 20) < 1 and abs(sy - 2) < 1 and sz < 30:
            drops.append(part)
        elif long > 2000 and min(sx, sy, sz) < 5:
            if sz < 10 and sx > 30:
                cleats.append(part)
            elif sy > 2000:
                joint_angles.append(part)
            else:
                cleats.append(part)
        elif max(sx, sz) < 100 and sy < 100:
            wall_brackets.append(part)
        else:
            joint_angles.append(part)

    return {
        "p": p,
        "g": g,
        "by": by,
        "slope_face": slope_face,
        "soffit_face": soffit_face,
        "trap_face": trap_face,
        "lid_area": lid_area,
        "slope_cd_lm": _length_m(by[LABEL_SLOPE_CD]),
        "soffit_cd_lm": _length_m(
            [c for c in by[LABEL_SOFFIT_CD] if c.bounding_box().size.Y > 1000.0]
        ),
        "soffit_rost_cd_lm": _length_m(
            [c for c in by[LABEL_SOFFIT_CD] if c.bounding_box().size.Y <= 1000.0]
        ),
        "soffit_ud_lm": _length_m(by[LABEL_SOFFIT_UD]),
        "n_soffit_rost_cd": len([c for c in by[LABEL_SOFFIT_CD] if c.bounding_box().size.Y <= 1000.0]),
        "bass_cd_lm": _length_m(by[LABEL_BASS_CD]),
        "strap_lm": _length_m(by[LABEL_RACKING_STRAP]),
        "duct_lm": _length_m(by[LABEL_SOFFIT_DUCT]),
        "n_direct": len(by[LABEL_SLOPE_DIRECT]),
        "n_slope_nonius": len(by[LABEL_SLOPE_NONIUS]),
        "n_soffit_nonius": len(nonius_hang),
        "n_drops": len(drops),
        "n_wall_brackets": len(wall_brackets),
        "joint_angle_lm": _length_m(joint_angles) if joint_angles else soffit_y / 1000.0,
        "cleat_lm": _length_m(cleats) if cleats else soffit_y / 1000.0,
        "n_bass_hanger": len(by[LABEL_BASS_HANGER]),
        "n_rafters": len(by[LABEL_RAFTERS]),
        "n_soffit_ud": len(by[LABEL_SOFFIT_UD]),
        "n_ducts": len(by[LABEL_SOFFIT_DUCT]),
        "plenum_wool_m3": _volume_m3(by[LABEL_PLENUM_WOOL]),
        "soffit_wool_m3": _volume_m3(by[LABEL_SOFFIT_WOOL]),
        "bass_wool_m3": _volume_m3(by[LABEL_BASS_WOOL]),
        "bass_gkb_m2": _volume_m3(by[LABEL_BASS_GKB]) / (p.bass_k_gkb / 1000.0),
    }


# Retail snapshots (Kč s DPH), Sep 2026 web checks — update as quotes arrive.
PRICE = {
    "nh80_m2": (534.0, f"přírodnístavba.cz · {URL_NATURHELD_80}"),
    "nh40_m2": (266.0, f"přírodnístavba.cz · {URL_NATURHELD_40}"),
    "rf_m2": (115.88, f"DEK · {URL_RF_DEK}"),
    "jutafol_m2": (35.73, f"IZOMAT JUTAFOL N AL 170 Speciál · {URL_JUTAFOL_SHOP}"),
    "gkb_m2": (96.0, "NonstopStavebniny / DEK RB~GKB orientačně"),
    "cd_bm": (29.21, f"DEK R-CD · {URL_CD_DEK}"),
    "direct_ks": (6.27, f"DEK KB510154 · {URL_DIRECT_DEK}"),
    "nonius_top_ks": (12.71, f"IZOMAT horní 340 · {URL_NONIUS_TOP_SHOP}"),
    "nonius_bot_ks": (10.43, f"DEK spodní CD · {URL_NONIUS_BOT_DEK}"),
    "nonius_pin_ks": (2.98, f"IZOMAT závlačka · {URL_NONIUS_PIN_SHOP}"),
    "trmen_ks": (477.0, "ceník Rigips KB517112 (sci-data / distributor) — ověřit nabídkou"),
    "strap_bm": (25.0, "orientační ocelová páska 40×2"),
    "duct_bm": (180.0, "orientační spirála Ø160 pozink"),
    "angle_bm": (55.0, "orientační L-profil / tenký úhelník"),
    "drop_ks": (8.0, "orientační krátký pásek/drop"),
    "bracket_ks": (35.0, "orientační nástěnný úhelník"),
}


def _nonius_set_price() -> tuple[float, str]:
    # 1 horní + 1 spodní + 2 závlačky (požární / vzpěr)
    top, st = PRICE["nonius_top_ks"]
    bot, sb = PRICE["nonius_bot_ks"]
    pin, sp = PRICE["nonius_pin_ks"]
    return top + bot + 2 * pin, f"sada: {st}; {sb}; 2× {sp}"


def build_lines(q) -> dict[str, list[Line]]:
    p = q["p"]
    face_s = q["slope_face"]
    face_k = q["soffit_face"]
    face_t = q["trap_face"]
    nonius_unit, nonius_src = _nonius_set_price()

    basic_kg_s = face_s * STOSILENT_BASIC_KG_M2
    finish_kg_s = face_s * STOSILENT_FINISH_KG_M2
    basic_kg_k = face_k * STOSILENT_BASIC_KG_M2
    finish_kg_k = face_k * STOSILENT_FINISH_KG_M2

    sikminy = [
        Line(
            1,
            "Povrchová úprava",
            "Akustický finální nátěr",
            _product("StoSilent Top Finish", URL_STOSILENT_FINISH),
            "spotřeba TL 3,0 kg/m² · balení 18 kg",
            finish_kg_s,
            "kg",
            10,
            None,
            f"TL Finish ≈3,0 kg/m² · {URL_STOSILENT_FINISH_TDS} · checklist {URL_STOSILENT_GUIDE}",
            f"Plocha {face_s:.2f} m² × 3,0 = {finish_kg_s:.1f} kg. "
            f"S rezervou ≈ {_fmt(finish_kg_s * 1.1, 1)} kg ≈ "
            f"{_ceil_pack(finish_kg_s * 1.1, STOSILENT_PAIL_KG)}× kbelík 18 kg. "
            "Cena u Sto CZ jen na poptávku.",
            True,
        ),
        Line(
            2,
            "Podkladní nátěr",
            "Akustický základní nátěr",
            _product("StoSilent Top Basic", URL_STOSILENT_BASIC),
            "spotřeba TL 2,5 kg/m² · balení 18 kg",
            basic_kg_s,
            "kg",
            10,
            None,
            f"Checklist StoSilent Distance: Basic 2,5 kg/m² · {URL_STOSILENT_GUIDE}",
            f"Plocha {face_s:.2f} m² × 2,5 = {basic_kg_s:.1f} kg. "
            f"S rezervou ≈ {_fmt(basic_kg_s * 1.1, 1)} kg ≈ "
            f"{_ceil_pack(basic_kg_s * 1.1, STOSILENT_PAIL_KG)}× kbelík 18 kg. "
            "Cena u Sto CZ jen na poptávku.",
            True,
        ),
        Line(
            3,
            "Akustická deska",
            "Dřevovláknitá deska ρ 140",
            _product("NaturHeld 140, tl. 80 mm (P+D 1880×615)", URL_NATURHELD_140),
            "80 mm",
            face_s,
            "m²",
            12,
            PRICE["nh80_m2"][0],
            PRICE["nh80_m2"][1],
            f"Šroubováno přes GKF do CD @ {p.cd_spacing:.0f} mm. Bez Flex / bez latí na šikmině.",
        ),
        Line(
            4,
            "Parozábrana",
            "Reflexní parozábrana Al",
            _product("Jutafol 145 Al (JUTA; retail N AL 170 Speciál 75 m²)", URL_JUTAFOL),
            "Al reflex · Sd vysoké · role 1,5×50 m",
            face_s,
            "m²",
            10,
            PRICE["jutafol_m2"][0],
            PRICE["jutafol_m2"][1],
            f"Na attic straně CD, pod minerální vlnou (stejně jako zbytek domu — "
            f"ne proměnlivá fólie). Plocha šikmin {face_s:.2f} m². "
            f"S rezervou 10 % ≈ {_fmt(face_s * 1.1, 1)} m² "
            f"≈ {_ceil_pack(face_s * 1.1, 75.0)}× role 75 m². "
            "Přesahy: JUTAFOL SP AL / SP 1.",
        ),
        Line(
            5,
            "Podkladní deska",
            "Sádrokarton GKF / RF",
            _product("Rigips RF (DF) 12,5 mm", URL_RF),
            f"{p.sdk_t} mm",
            face_s,
            "m²",
            12,
            PRICE["rf_m2"][0],
            PRICE["rf_m2"][1],
            "Končí u zlomu na víko kastlíku.",
        ),
        Line(
            6,
            "Rošt CD",
            "Profil CD 60×27",
            _product("Rigips R-CD profil 27/60/27", URL_CD),
            f"60×27 mm @ {p.cd_spacing:.0f} mm ⊥ krokvím",
            q["slope_cd_lm"],
            "bm",
            10,
            PRICE["cd_bm"][0],
            PRICE["cd_bm"][1],
            f"{len(q['by'][LABEL_SLOPE_CD])} profilů × {p.room_length/1000:.2f} m.",
        ),
        Line(
            7,
            "Závěs — okenní šikmina",
            "Přímý závěs 125 mm",
            _product("Rigips přímý závěs 125 mm, KB510154", URL_DIRECT),
            "125 mm",
            q["n_direct"],
            "ks",
            8,
            PRICE["direct_ks"][0],
            PRICE["direct_ks"][1],
            "Kotvení z boku krokve.",
        ),
        Line(
            8,
            "Závěs — šikmina u skříní",
            "Nonius sada (horní+spodní+2× závlačka)",
            _product("Rigips Nonius horní 340/440 + spodní CD + 2× závlačka", URL_NONIUS_TOP),
            "horní 340/440 dle svěšení",
            q["n_slope_nonius"],
            "ks",
            8,
            nonius_unit,
            nonius_src,
            "Svěšení otevírá se k nábytku.",
        ),
        Line(
            9,
            "Výplň dutiny / mezi krokvemi",
            "Minerální vlna",
            "Výrobek dosud nezvolen",
            "void ~80 mm + mezi krokvemi 160 mm",
            q["plenum_wool_m3"],
            "m³",
            10,
            None,
            "—",
            "Nedořešeno: typ vlny (akustika σ≈12000). Ne NaturHeld Flex.",
            True,
        ),
        Line(
            10,
            "Zavětrování",
            "Ocelová páska",
            "Výrobek dosud nezvolen",
            f"{p.strap_w:.0f}×{p.strap_t:.0f} mm @ 45°",
            q["strap_lm"],
            "bm",
            10,
            PRICE["strap_bm"][0],
            PRICE["strap_bm"][1],
            "Typ pásky / kotevní prvky ještě sjednotit.",
            True,
        ),
        Line(
            11,
            "Krov (nosná konstrukce)",
            "Krokve",
            "Stávající krov",
            f"{p.rafter_w:.0f}/{p.rafter_t:.0f} @ {p.rafter_spacing:.0f}",
            q["n_rafters"],
            "ks",
            None,
            None,
            "—",
            "Není položka nákupu akustiky.",
        ),
    ]

    kastlik = [
        Line(
            1,
            "Povrchová úprava (bok + spodní líce)",
            "Akustický finální nátěr",
            _product("StoSilent Top Finish", URL_STOSILENT_FINISH),
            "3,0 kg/m² · 18 kg",
            finish_kg_k,
            "kg",
            10,
            None,
            f"TL Finish 3,0 kg/m² · {URL_STOSILENT_FINISH_TDS}",
            f"{face_k:.2f} m² × 3,0 = {finish_kg_k:.1f} kg; "
            f"s rezervou ≈ {_ceil_pack(finish_kg_k * 1.1, STOSILENT_PAIL_KG)}× 18 kg. Cena na poptávku.",
            True,
        ),
        Line(
            2,
            "Podkladní nátěr",
            "Akustický základní nátěr",
            _product("StoSilent Top Basic", URL_STOSILENT_BASIC),
            "2,5 kg/m² · 18 kg",
            basic_kg_k,
            "kg",
            10,
            None,
            f"Checklist 2,5 kg/m² · {URL_STOSILENT_GUIDE}",
            f"{face_k:.2f} m² × 2,5 = {basic_kg_k:.1f} kg; "
            f"s rezervou ≈ {_ceil_pack(basic_kg_k * 1.1, STOSILENT_PAIL_KG)}× 18 kg. Cena na poptávku.",
            True,
        ),
        Line(
            3,
            "Akustická deska L",
            "Dřevovláknitá deska ρ 140",
            _product("NaturHeld 140, tl. 40 mm", URL_NATURHELD_140),
            "40 mm",
            face_k,
            "m²",
            12,
            PRICE["nh40_m2"][0],
            PRICE["nh40_m2"][1],
            "Butuje na X_FURN; mezera ke skříním 20 mm.",
        ),
        Line(
            4,
            "Výplň akustické dutiny",
            "Minerální vlna",
            "Výrobek dosud nezvolen",
            "za NH, zastaveno před potrubím",
            q["soffit_wool_m3"],
            "m³",
            10,
            None,
            "—",
            "Nedořešeno: typ vlny.",
            True,
        ),
        Line(
            5,
            "Rošt kastlíku",
            "Profil CD 60×27",
            _product("Rigips R-CD 27/60/27", URL_CD),
            f"60×27 @ {p.rost_spacing:.0f}",
            q["soffit_rost_cd_lm"],
            "bm",
            10,
            PRICE["cd_bm"][0],
            PRICE["cd_bm"][1],
            f"{q['n_soffit_rost_cd']} prvků (svislé + spodní). "
            f"UD {p.ud_w:.0f}×{p.cd_t:.0f} u zdi {q['soffit_ud_lm']:.2f} bm "
            f"({q['n_soffit_ud']} ks). Žádné latě, žádné podložky.",
        ),
        Line(
            6,
            "Víko kastlíku",
            "Sádrokarton GKF / RF",
            _product("Rigips RF (DF) 12,5 mm", URL_RF),
            f"{p.sdk_t} mm",
            q["lid_area"],
            "m²",
            12,
            PRICE["rf_m2"][0],
            PRICE["rf_m2"][1],
            "Horní líce v úrovni pozednice.",
        ),
        Line(
            7,
            "Vodorovný CD",
            "Profil CD 60×27",
            _product("Rigips R-CD", URL_CD),
            "60×27 mm",
            q["soffit_cd_lm"],
            "bm",
            10,
            PRICE["cd_bm"][0],
            PRICE["cd_bm"][1],
            "Jedna řada = jediný Nonius řádek.",
        ),
        Line(
            8,
            "Závěs CD → krokve",
            "Nonius sada",
            _product("Rigips Nonius sada", URL_NONIUS_TOP),
            "≈400 mm svěšení",
            q["n_soffit_nonius"],
            "ks",
            8,
            nonius_unit,
            nonius_src,
            "1 závěs na krokev na rostovém CD.",
        ),
        Line(
            9,
            "Drop CD → svislý CD",
            "Krátký ocelový závěs",
            "Výrobek dosud nezvolen",
            f"{p.soffit_drop_w:.0f}×{p.soffit_drop_t:.0f} mm",
            q["n_drops"],
            "ks",
            10,
            PRICE["drop_ks"][0],
            PRICE["drop_ks"][1],
            "Nedořešeno: katalog vs. pásek na míru.",
            True,
        ),
        Line(
            10,
            "Úhelník skrytého spoje",
            "Tenkostěnný L-profil",
            "Výrobek dosud nezvolen",
            "podél spoje šikmina/víko",
            q["joint_angle_lm"],
            "bm",
            10,
            PRICE["angle_bm"][0],
            PRICE["angle_bm"][1],
            "Nedořešeno: přesný profil.",
            True,
        ),
        Line(
            11,
            "Vzpěra u obvodové zdi",
            "Nástěnný úhelník",
            "Výrobek dosud nezvolen",
            f"rameno ≈ {p.wall_bracket_leg:.0f} mm",
            q["n_wall_brackets"],
            "ks",
            10,
            PRICE["bracket_ks"][0],
            PRICE["bracket_ks"][1],
            "Jen vzpěra — ne závěsný bod.",
            True,
        ),
        Line(
            12,
            "Cleat víka na pozednici",
            "Tenkostěnný úhelník",
            "Výrobek dosud nezvolen",
            "podél líce pozednice",
            q["cleat_lm"],
            "bm",
            10,
            PRICE["angle_bm"][0],
            PRICE["angle_bm"][1],
            "Pozednice nenosí kastlík.",
            True,
        ),
        Line(
            13,
            "VZT / klimatizace",
            "Spirála Ø160",
            "Spirála Ø160 (3 trasy)",
            f"OD {p.duct_od:.0f} mm",
            q["duct_lm"],
            "bm",
            8,
            PRICE["duct_bm"][0],
            PRICE["duct_bm"][1],
            f"{q['n_ducts']} trubky; výrobce VZT doplnit.",
            True,
        ),
    ]

    stity = [
        Line(
            "A1",
            "Lícová deska (membrána)",
            "Sádrokarton GKB / RF plná",
            _product("Rigips RB (A) / RF 12,5 — plná, NE perforovaná", URL_GKB),
            f"{p.bass_k_gkb} mm",
            q["bass_gkb_m2"],
            "m²",
            12,
            PRICE["gkb_m2"][0],
            PRICE["gkb_m2"][1],
            "Včetně spodního víka. Soft-joint k NH šikminy.",
        ),
        Line(
            "A2",
            "Rám CD / UW",
            "Profil CD 60×27",
            _product("Rigips R-CD 27/60/27", URL_CD),
            "kuchyň 625 svisle + dořez; obývák 1000 + dořez",
            q["bass_cd_lm"],
            "bm",
            10,
            PRICE["cd_bm"][0],
            PRICE["cd_bm"][1],
            "Lišty pod spárou: kuchyň 2000 mm, obývák 1250 a 2500 mm, na CD. UW v modelu jako CD.",
        ),
        Line(
            "A3",
            "Kotvení ke zdi (vzadu)",
            "Akustický třmen se Sylomerem",
            _product("Rigips KB517112", URL_TRMEN),
            "dosah ≤ ~120 mm",
            q["n_bass_hanger"],
            "ks",
            8,
            PRICE["trmen_ks"][0],
            PRICE["trmen_ks"][1],
            "Jen k zadnímu CD. Alt. Knauf W623.",
        ),
        Line(
            "A4",
            "Závěs shora",
            "Nonius / akust. stropní závěs",
            _product("Rigips Nonius / akust. závěs", URL_NONIUS_TOP),
            "k hornímu UW/CD",
            0,
            "ks",
            None,
            None,
            "—",
            "Nedořešeno v 3D počtu — hybrid dle detailů; model kotví hlavně třmeny ke štítu.",
            True,
        ),
        Line(
            "B1",
            "MW kuchyně u zdi",
            "Minerální vlna",
            "Výrobek dosud nezvolen",
            f"{p.bass_k_wool} mm",
            face_t * p.bass_k_wool / 1000.0,
            "m³",
            10,
            None,
            "—",
            f"Split {p.bass_k_wool}+{p.bass_k_air}+{p.bass_k_gkb} = {p.predstena_kitchen:.0f} mm.",
            True,
        ),
        Line(
            "C2",
            "MW obývák u zdi",
            "Minerální vlna",
            "Výrobek dosud nezvolen",
            f"{p.bass_l_wool} mm",
            face_t * p.bass_l_wool / 1000.0,
            "m³",
            10,
            None,
            "—",
            f"Split {p.bass_l_gkb}+{p.bass_l_air}+{p.bass_l_wool} = {p.predstena_living:.0f} mm.",
            True,
        ),
    ]

    # Rolled-up souhrn (unique materials)
    basic_kg = basic_kg_s + basic_kg_k
    finish_kg = finish_kg_s + finish_kg_k
    nonius_ks = q["n_slope_nonius"] + q["n_soffit_nonius"]
    wool = q["plenum_wool_m3"] + q["soffit_wool_m3"] + q["bass_wool_m3"]
    gkf = face_s + q["lid_area"]
    cd = q["slope_cd_lm"] + q["soffit_cd_lm"] + q["soffit_rost_cd_lm"] + q["bass_cd_lm"]

    souhrn = [
        Line(
            1,
            "StoSilent Top Finish",
            "Akustický finál",
            _product("StoSilent Top Finish", URL_STOSILENT_FINISH),
            "3,0 kg/m² · 18 kg",
            finish_kg,
            "kg",
            10,
            None,
            f"TL 3,0 kg/m² · {URL_STOSILENT_FINISH_TDS}",
            f"Šikminy+kastlík {(face_s+face_k):.2f} m² → {finish_kg:.1f} kg; "
            f"objednat {_ceil_pack(finish_kg * 1.1, STOSILENT_PAIL_KG)}× 18 kg. Cena na poptávku Sto.",
            True,
        ),
        Line(
            2,
            "StoSilent Top Basic",
            "Akustický základ",
            _product("StoSilent Top Basic", URL_STOSILENT_BASIC),
            "2,5 kg/m² · 18 kg",
            basic_kg,
            "kg",
            10,
            None,
            f"Checklist 2,5 kg/m² · {URL_STOSILENT_GUIDE}",
            f"{basic_kg:.1f} kg; objednat {_ceil_pack(basic_kg * 1.1, STOSILENT_PAIL_KG)}× 18 kg. Cena na poptávku Sto.",
            True,
        ),
        Line(
            3,
            "NaturHeld 140 — 80 mm",
            "Dřevovláknitá deska",
            _product("NaturHeld 140 / 80 mm", URL_NATURHELD_140),
            "80 mm",
            face_s,
            "m²",
            12,
            PRICE["nh80_m2"][0],
            PRICE["nh80_m2"][1],
            "Šikminy.",
        ),
        Line(
            4,
            "NaturHeld 140 — 40 mm",
            "Dřevovláknitá deska",
            _product("NaturHeld 140 / 40 mm", URL_NATURHELD_140),
            "40 mm",
            face_k,
            "m²",
            12,
            PRICE["nh40_m2"][0],
            PRICE["nh40_m2"][1],
            "Kastlík.",
        ),
        Line(
            5,
            "Parozábrana",
            "Jutafol 145 Al",
            _product("Jutafol 145 Al", URL_JUTAFOL),
            "role 75 m²",
            face_s,
            "m²",
            10,
            PRICE["jutafol_m2"][0],
            PRICE["jutafol_m2"][1],
            f"Šikminy {face_s:.2f} m² (+10 % ≈ {_fmt(face_s * 1.1, 1)} m²). Nad CD, pod vlnou.",
        ),
        Line(
            6,
            "Sádrokarton RF 12,5",
            "GKF/RF",
            _product("Rigips RF 12,5", URL_RF),
            "12,5 mm",
            gkf,
            "m²",
            12,
            PRICE["rf_m2"][0],
            PRICE["rf_m2"][1],
            "Šikminy + víko kastlíku.",
        ),
        Line(
            7,
            "Sádrokarton GKB/RF 12,5 (basstrapy)",
            "GKB plná",
            _product("Rigips RB/RF 12,5", URL_GKB),
            "12,5 mm",
            q["bass_gkb_m2"],
            "m²",
            12,
            PRICE["gkb_m2"][0],
            PRICE["gkb_m2"][1],
            "Štíty.",
        ),
        Line(
            8,
            "Profil R-CD 60×27",
            "CD profil",
            _product("Rigips R-CD", URL_CD),
            "60×27",
            cd,
            "bm",
            10,
            PRICE["cd_bm"][0],
            PRICE["cd_bm"][1],
            "Šikminy + kastlík + štíty.",
        ),
        Line(
            9,
            "Přímý závěs KB510154",
            "Přímý závěs 125",
            _product("Rigips KB510154", URL_DIRECT),
            "125 mm",
            q["n_direct"],
            "ks",
            8,
            PRICE["direct_ks"][0],
            PRICE["direct_ks"][1],
            "Okenní šikmina.",
        ),
        Line(
            10,
            "Nonius sada",
            "H+S+2× závlačka",
            _product("Rigips Nonius sada", URL_NONIUS_TOP),
            "340/440",
            nonius_ks,
            "ks",
            8,
            nonius_unit,
            nonius_src,
            "Šikminy u skříní + kastlík.",
        ),
        Line(
            11,
            "Akustický třmen KB517112",
            "Třmen + Sylomer",
            _product("Rigips KB517112", URL_TRMEN),
            "—",
            q["n_bass_hanger"],
            "ks",
            8,
            PRICE["trmen_ks"][0],
            PRICE["trmen_ks"][1],
            "Štíty.",
        ),
        Line(
            12,
            "Minerální vlna (souhrn)",
            "Minerální vlna",
            "Výrobek dosud nezvolen",
            "—",
            wool,
            "m³",
            10,
            None,
            "—",
            f"Plenum {q['plenum_wool_m3']:.2f} + kastlík {q['soffit_wool_m3']:.2f} "
            f"+ štíty {q['bass_wool_m3']:.2f} m³.",
            True,
        ),
        Line(
            13,
            "UD 28×27",
            "UD profil",
            "Rigips UD (obvod kastlíku)",
            f"{p.ud_w:.0f}×{p.cd_t:.0f}",
            q["soffit_ud_lm"],
            "bm",
            10,
            PRICE["cd_bm"][0],
            PRICE["cd_bm"][1],
            "U obvodové zdi kastlíku. Cena orientačně jako R-CD — UD ověřit nabídkou.",
            True,
        ),
        Line(
            14,
            "Zavětrovací pásky 40×2",
            "Ocelová páska",
            "Výrobek dosud nezvolen",
            "40×2",
            q["strap_lm"],
            "bm",
            10,
            PRICE["strap_bm"][0],
            PRICE["strap_bm"][1],
            "Nedořešeno: přesný typ.",
            True,
        ),
        Line(
            15,
            "Spirála Ø160",
            "VZT potrubí",
            "Výrobce VZT dosud nezvolen",
            "Ø160",
            q["duct_lm"],
            "bm",
            8,
            PRICE["duct_bm"][0],
            PRICE["duct_bm"][1],
            "3 trasy v kastlíku.",
            True,
        ),
        Line(
            16,
            "Dropy kastlíku",
            "Krátký závěs",
            "Dosud nezvolen",
            "20×2",
            q["n_drops"],
            "ks",
            10,
            PRICE["drop_ks"][0],
            PRICE["drop_ks"][1],
            "Kastlík.",
            True,
        ),
        Line(
            17,
            "Úhelníky kastlíku (spoj + cleat)",
            "L-profil",
            "Dosud nezvolen",
            "—",
            q["joint_angle_lm"] + q["cleat_lm"],
            "bm",
            10,
            PRICE["angle_bm"][0],
            PRICE["angle_bm"][1],
            "Spoj + cleat na pozednici.",
            True,
        ),
        Line(
            18,
            "Nástěnné úhelníky kastlíku",
            "Úhelník",
            "Dosud nezvolen",
            "—",
            q["n_wall_brackets"],
            "ks",
            10,
            PRICE["bracket_ks"][0],
            PRICE["bracket_ks"][1],
            "Vzpěra u zdi.",
            True,
        ),
    ]

    return {"Šikminy": sikminy, "Kastlík": kastlik, "Štíty": stity, "Souhrn materiálu": souhrn}


def _style_header(ws):
    for col, title in enumerate(COLS, 1):
        cell = ws.cell(1, col, title)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"
    widths = [8, 26, 22, 42, 22, 12, 9, 12, 14, 14, 14, 36, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _section(ws, row: int, title: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLS))
    for col in range(1, len(COLS) + 1):
        cell = ws.cell(row, col)
        cell.fill = SECTION_FILL
        cell.border = THIN
    cell = ws.cell(row, 1, title)
    cell.font = SECTION_FONT
    cell.alignment = WRAP
    ws.row_dimensions[row].height = 22
    return row + 1


def _write_line(ws, row: int, line: Line) -> int:
    reserve = "" if line.reserve_pct is None else f"{_fmt(line.reserve_pct)} %"
    qty_r = _fmt(line.qty_with_reserve, 2 if line.unit in {"kg", "m³", "m²"} else 1)
    price = "" if line.unit_price is None else round(line.unit_price, 2)
    total = "" if line.line_total is None else round(line.line_total, 0)
    values = [
        line.order,
        line.layer,
        line.product_type,
        line.product,
        line.dimension,
        _fmt(line.qty_net, 2 if line.unit in {"kg", "m³", "m²"} else 1),
        line.unit,
        reserve,
        qty_r,
        price,
        total,
        line.price_source,
        line.note,
    ]
    for col, val in enumerate(values, 1):
        cell = ws.cell(row, col, val if val != "" else None)
        cell.alignment = MONEY if col in (6, 8, 9, 10, 11) else WRAP
        cell.border = THIN
        if line.open_item:
            cell.fill = OPEN_FILL
        if col == 4 and isinstance(val, str) and "\n" in val:
            name, url = val.split("\n", 1)
            cell.value = name
            if url.startswith("http"):
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single", name="Calibri", size=11)
        if col in (10, 11) and isinstance(cell.value, (int, float)):
            cell.number_format = "#,##0.00" if col == 10 else "#,##0"
    ws.row_dimensions[row].height = 52 if line.open_item else 44
    return row + 1


def _write_totals(ws, row: int, lines: list[Line]) -> int:
    priced = [ln.line_total for ln in lines if ln.line_total is not None]
    missing = [ln.layer for ln in lines if ln.unit_price is None and ln.qty_net]
    total = sum(priced)
    row = _section(ws, row, "")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    cell = ws.cell(row, 1, "SOUČET OCENĚNÝCH POLOŽEK (s rezervou, Kč s DPH)")
    cell.fill = TOTAL_FILL
    cell.font = Font(bold=True, name="Calibri", size=11)
    for col in range(1, len(COLS) + 1):
        ws.cell(row, col).fill = TOTAL_FILL
        ws.cell(row, col).border = THIN
    tot_cell = ws.cell(row, 11, round(total, 0))
    tot_cell.number_format = "#,##0"
    tot_cell.font = Font(bold=True, name="Calibri", size=12)
    tot_cell.fill = TOTAL_FILL
    tot_cell.alignment = MONEY
    ws.row_dimensions[row].height = 24
    row += 1
    if missing:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLS))
        cell = ws.cell(
            row,
            1,
            "Neoceněno (chybí v součtu): " + "; ".join(missing),
        )
        cell.fill = NOTE_FILL
        cell.alignment = WRAP
        for col in range(1, len(COLS) + 1):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).fill = NOTE_FILL
        ws.row_dimensions[row].height = 36
        row += 1
    return row


def write_assembly_sheet(wb, title: str, lines: list[Line], intro: str, index: int):
    ws = wb.create_sheet(title, index)
    _style_header(ws)
    row = 2
    row = _section(ws, row, intro)
    for line in lines:
        row = _write_line(ws, row, line)
    _write_totals(ws, row, lines)
    return ws


def write_meta(wb, q, index: int):
    ws = wb.create_sheet("Meta", index)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Položka"
    ws["B1"] = "Hodnota"
    for col in (1, 2):
        ws.cell(1, col).fill = HDR_FILL
        ws.cell(1, col).font = HDR_FONT
    p = q["p"]
    rows = [
        ("Model", "obyvak (Obývák 1.02)"),
        ("Regenerace", "PYTHONPATH=models:src python scripts/export_obyvak_skladba.py"),
        ("StoSilent Basic", f"{STOSILENT_BASIC_KG_M2} kg/m² (StoSilent Distance checklist / TL)"),
        ("StoSilent Finish", f"{STOSILENT_FINISH_KG_M2} kg/m² (TL Top Finish + checklist)"),
        ("StoSilent balení", f"{STOSILENT_PAIL_KG:.0f} kg / kbelík"),
        ("Ceny", "Kč s DPH, webové snapshoty — ověřit aktuální nabídkou; Sto jen na poptávku"),
        ("Rezerva desky/SDK", "12 % (prořez)"),
        ("Rezerva CD / závěsy", "8–10 %"),
        ("Rezerva nátěry / vlna", "10 %"),
        ("Plocha šikmin", f"{q['slope_face']:.2f} m²"),
        ("Plocha kastlíku L", f"{q['soffit_face']:.2f} m²"),
        ("Plocha jednoho štítu", f"{q['trap_face']:.2f} m²"),
        ("Místnost", f"{p.room_width:.0f} × {p.room_length:.0f} mm, sklon {p.roof_angle_deg:.0f}°"),
        (
            "iOS / Numbers",
            "Otevři docs/skladba.html na Pages preview — tlačítka stáhnou .xlsx / .csv. "
            "Chat nepřiloží xlsx (jen obrázky/video).",
        ),
    ]
    for i, (a, b) in enumerate(rows, 2):
        ws.cell(i, 1, a).border = THIN
        ws.cell(i, 2, b).border = THIN
        ws.cell(i, 2).alignment = WRAP
        ws.row_dimensions[i].height = 28
    return ws


def write_csv(lines: list[Line]) -> None:
    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(COLS)
        for line in lines:
            reserve = "" if line.reserve_pct is None else f"{_fmt(line.reserve_pct)} %"
            prod = line.product.replace("\n", " | ")
            w.writerow(
                [
                    line.order,
                    line.layer,
                    line.product_type,
                    prod,
                    line.dimension,
                    _fmt(line.qty_net, 2),
                    line.unit,
                    reserve,
                    _fmt(line.qty_with_reserve, 2),
                    "" if line.unit_price is None else round(line.unit_price, 2),
                    "" if line.line_total is None else round(line.line_total, 0),
                    line.price_source,
                    line.note,
                ]
            )
        priced = sum(ln.line_total for ln in lines if ln.line_total is not None)
        w.writerow([])
        w.writerow(["", "SOUČET OCENĚNÝCH (s rezervou)", "", "", "", "", "", "", "", "", round(priced, 0)])


def write_html() -> None:
    OUT_HTML.write_text(
        """<!DOCTYPE html>
<html lang="cs">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Obývák — skladba materiálu</title>
  <style>
    :root { color-scheme: light; }
    body {
      margin: 0; font-family: "Avenir Next", "Segoe UI", sans-serif;
      background: linear-gradient(160deg, #e8eef5 0%, #f7f3ea 55%, #f0e6d8 100%);
      color: #1a2433; min-height: 100vh;
    }
    main { max-width: 34rem; margin: 0 auto; padding: 2.5rem 1.25rem 4rem; }
    h1 { font-size: 1.75rem; letter-spacing: -0.02em; margin: 0 0 0.35rem; }
    p { line-height: 1.45; margin: 0 0 1rem; color: #334155; }
    .actions { display: grid; gap: 0.75rem; margin: 1.5rem 0; }
    a.btn {
      display: block; text-align: center; text-decoration: none;
      padding: 0.95rem 1.1rem; border-radius: 0.65rem;
      background: #1f4e79; color: #fff; font-weight: 600;
    }
    a.btn.secondary { background: #0f766e; }
    .hint {
      font-size: 0.92rem; background: rgba(255,255,255,0.7);
      border: 1px solid #d5dde8; border-radius: 0.65rem; padding: 0.9rem 1rem;
    }
  </style>
</head>
<body>
  <main>
    <h1>Obývák — skladba</h1>
    <p>Šikminy, kastlík a štíty. Pro Numbers na iPhonu/iPadu stačí klepnout — Safari nabídne otevření v Numbers.</p>
    <div class="actions">
      <a class="btn" href="obyvak-skladba-materialu.xlsx">Stáhnout Excel (.xlsx)</a>
      <a class="btn secondary" href="obyvak-skladba-souhrn.csv">Stáhnout souhrn CSV (Numbers)</a>
    </div>
    <p class="hint">CSV je nejjednodušší cesta do Numbers. XLSX otevře Numbers přímo z prohlížeče. Soubory jsou na GitHub Pages u tohoto PR / na produkci po merge.</p>
  </main>
</body>
</html>
""",
        encoding="utf-8",
    )


def main() -> int:
    print("Counting solids from obyvak build…")
    q = collect_quantities()
    groups = build_lines(q)

    wb = Workbook()
    wb.remove(wb.active)
    write_assembly_sheet(
        wb,
        "Šikminy",
        groups["Šikminy"],
        "Skladba interiér → půda. StoSilent kg = plocha × TL (Basic 2,5 / Finish 3,0 kg/m²).",
        0,
    )
    write_assembly_sheet(
        wb,
        "Kastlík",
        groups["Kastlík"],
        "Kastlík nad skříněmi — závěs z CD, ne z nábytku. StoSilent kg z TL.",
        1,
    )
    write_assembly_sheet(
        wb,
        "Štíty",
        groups["Štíty"],
        "Basstrapy kuchyně 190 / obývák 450 pod kontinuální šikminou.",
        2,
    )
    write_assembly_sheet(
        wb,
        "Souhrn materiálu",
        groups["Souhrn materiálu"],
        "Srolované množství napříč sestavami · součty cen s rezervou (Kč s DPH).",
        3,
    )
    write_meta(wb, q, 4)

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    write_csv(groups["Souhrn materiálu"])
    write_html()

    priced = sum(ln.line_total for ln in groups["Souhrn materiálu"] if ln.line_total is not None)
    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_CSV}")
    print(f"Wrote {OUT_HTML}")
    print(
        f"Areas m²: slopes={q['slope_face']:.2f} soffit={q['soffit_face']:.2f} "
        f"trap={q['trap_face']:.2f}"
    )
    print(f"StoSilent Basic {q['slope_face']*STOSILENT_BASIC_KG_M2 + q['soffit_face']*STOSILENT_BASIC_KG_M2:.1f} kg")
    print(f"StoSilent Finish {q['slope_face']*STOSILENT_FINISH_KG_M2 + q['soffit_face']*STOSILENT_FINISH_KG_M2:.1f} kg")
    print(f"Priced subtotal (with reserve): {priced:,.0f} Kč")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
